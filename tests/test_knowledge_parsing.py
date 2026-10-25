"""knowledge_base 的解析/检索路径测试（不加载真实 embedding）。"""

from __future__ import annotations

import json
import sys
from dataclasses import MISSING, fields
from pathlib import Path
from types import SimpleNamespace

import docx
import pytest
from langchain.tools import ToolRuntime
from langchain_core.documents import Document

from personal_assistant import knowledge_base as kb
from personal_assistant.business_tools import RunContext


def _runtime(role: str = "employee") -> ToolRuntime:
    kwargs: dict[str, object] = {}
    for field_info in fields(ToolRuntime):
        if field_info.name == "context":
            kwargs["context"] = RunContext(
                user_id="USR-001",
                username="zhang_wei",
                display_name="张伟",
                role=role,
                department_id="DEPT-001",
            )
        elif field_info.name == "state":
            kwargs["state"] = {}
        elif field_info.name == "tools":
            kwargs["tools"] = []
        elif field_info.default is MISSING and field_info.default_factory is MISSING:
            kwargs[field_info.name] = None
    return ToolRuntime(**kwargs)


def test_text_helpers(tmp_path: Path) -> None:
    path = tmp_path / "doc.md"
    path.write_text("﻿# 标题\n\n\n正文  ", encoding="utf-8")

    text = kb._read_text(path)
    assert text.startswith("# 标题")
    assert kb._collapse_blank_lines("a\r\n\r\n\r\nb\r\n") == "a\n\nb"
    assert kb._ensure_nonempty("x", path) == "x"
    with pytest.raises(ValueError, match="解析结果为空"):
        kb._ensure_nonempty("  ", path)


def test_parse_text_and_empty(tmp_path: Path) -> None:
    path = tmp_path / "doc.txt"
    path.write_text("hello", encoding="utf-8")
    parsed = kb._parse_text(path)
    assert parsed.parser_mode_actual == "text"
    assert parsed.blocks[0].text == "hello"

    path.write_text("   ", encoding="utf-8")
    with pytest.raises(ValueError, match="解析结果为空"):
        kb._parse_text(path)


def test_parse_html_strips_script_and_keeps_title(tmp_path: Path) -> None:
    path = tmp_path / "doc.html"
    path.write_text(
        "<html><head><title>标题</title><style>x{}</style></head>"
        "<body><script>bad()</script><p>正文</p></body></html>",
        encoding="utf-8",
    )

    parsed = kb._parse_html(path)

    assert parsed.parser_mode_actual == "html_text"
    assert "标题" in parsed.blocks[0].text
    assert "正文" in parsed.blocks[0].text
    assert "bad()" not in parsed.blocks[0].text


def test_parse_docx_paragraphs_and_table(tmp_path: Path) -> None:
    path = tmp_path / "doc.docx"
    document = docx.Document()
    document.add_paragraph("段落一")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "列1"
    table.cell(0, 1).text = "列2"
    document.save(path)

    parsed = kb._parse_docx(path)

    assert parsed.parser_mode_actual == "document_text"
    assert any(block.kind == "table" for block in parsed.blocks)
    assert any("列1" in block.text for block in parsed.blocks)


def test_parse_pdf_text(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    class _Page:
        def extract_text(self) -> str:
            return "第一页\r\n\r\n第二行"

    class _Reader:
        pages = [_Page()]

    import pypdf

    monkeypatch.setattr(pypdf, "PdfReader", lambda _path: _Reader())
    path = tmp_path / "doc.pdf"
    path.write_bytes(b"%PDF")

    parsed = kb._parse_pdf_text(path)

    assert parsed.parser_mode_actual == "pdf_text"
    assert "第一页" in parsed.blocks[0].text


def test_ocr_lines_and_image_ocr(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    entries = [
        [[[0, 0], [10, 0], [10, 10], [0, 10]], "第一行", 0.9],
        [[[0, 20], [10, 20], [10, 30], [0, 30]], "第二行", 0.9],
    ]
    assert kb._ocr_lines(entries) == ["第一行", "第二行"]
    assert kb._ocr_lines([]) == []
    assert kb._ocr_lines([["bad"]]) == []

    path = tmp_path / "img.png"
    path.write_bytes(b"x")
    monkeypatch.setattr(kb, "_run_ocr", lambda _image: ["文字"])
    parsed = kb._parse_image_ocr(path)
    assert parsed.parser_mode_actual == "ocr"
    assert parsed.is_machine_extracted is True

    monkeypatch.setattr(kb, "_run_ocr", lambda _image: [])
    with pytest.raises(ValueError, match="OCR 结果为空"):
        kb._parse_image_ocr(path)


def test_parse_ocr_pdf(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    class _Rendered:
        def to_pil(self) -> object:
            return object()

    class _Page:
        def render(self, scale: float) -> _Rendered:
            return _Rendered()

    class _Pdf:
        def __iter__(self):
            return iter([_Page()])

        def close(self) -> None:
            return None

    fake_module = SimpleNamespace(PdfDocument=lambda _path: _Pdf())
    monkeypatch.setitem(sys.modules, "pypdfium2", fake_module)
    monkeypatch.setattr(kb, "_run_ocr", lambda _image: ["OCR 行"])
    path = tmp_path / "scan.pdf"
    path.write_bytes(b"%PDF")

    parsed = kb._parse_ocr_pdf(path)

    assert parsed.parser_mode_actual == "ocr"
    assert "OCR 行" in parsed.blocks[0].text


def test_describe_image_and_fallback(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    class _Model:
        def invoke(self, _messages: object) -> SimpleNamespace:
            return SimpleNamespace(content="视觉描述")

    monkeypatch.setattr(
        kb,
        "build_chat_model",
        lambda _prefix: (_Model(), "vision-model", "vision"),
    )
    path = tmp_path / "img.png"
    path.write_bytes(b"x")

    parsed = kb._describe_image(path, "prompt", "vision_multimodal")
    assert parsed.parser_mode_actual == "vision_multimodal"
    assert parsed.provenance["model_name"] == "vision-model"

    monkeypatch.setattr(
        kb,
        "_describe_image",
        lambda *_args: (_ for _ in ()).throw(RuntimeError("vision down")),
    )
    monkeypatch.setattr(
        kb,
        "_parse_image_ocr",
        lambda _path: kb.ParsedDocument(
            blocks=[kb.Block("ocr")], parser_mode_actual="ocr"
        ),
    )
    fallback = kb._vision_or_ocr_fallback(path, "prompt")
    assert fallback.parser_mode_actual == "ocr_fallback"
    assert "vision down" in fallback.provenance["vision_error"]


def test_search_knowledge_paths(monkeypatch: pytest.MonkeyPatch) -> None:
    class _Store:
        def similarity_search_with_score(self, *args: object, **kwargs: object):
            return [
                (
                    Document(
                        page_content="年假规定",
                        metadata={
                            "document_id": "KB-PUB-001",
                            "title": "员工手册",
                            "path": "public/handbook.md",
                            "chunk_id": "KB-PUB-001#001",
                            "access_scope": "public",
                        },
                    ),
                    0.5,
                )
            ]

    monkeypatch.setattr(kb, "_open_vector_store", lambda: _Store())
    runtime = _runtime()

    payload = json.loads(
        kb.search_knowledge.invoke({"query": "年假", "runtime": runtime})
    )
    assert payload["ok"] is True
    assert payload["data"]["count"] == 1

    empty = json.loads(kb.search_knowledge.invoke({"query": "", "runtime": runtime}))
    assert empty["error"] == "invalid_argument"

    denied = json.loads(
        kb.search_knowledge.invoke({"query": "年假", "runtime": _runtime(role="guest")})
    )
    assert denied["error"] == "permission_denied"


def test_search_knowledge_no_evidence(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        kb,
        "_open_vector_store",
        lambda: SimpleNamespace(similarity_search_with_score=lambda *a, **k: []),
    )
    payload = json.loads(
        kb.search_knowledge.invoke({"query": "年假", "runtime": _runtime()})
    )
    assert payload["error"] == "no_evidence"


def test_build_vector_store_with_fakes(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    knowledge = tmp_path / "knowledge"
    knowledge.mkdir()
    (knowledge / "doc.md").write_text("第一段。第二段。", encoding="utf-8")
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "documents": [
                    {
                        "document_id": "KB-PUB-001",
                        "title": "文档",
                        "path": "doc.md",
                        "parser_mode": "text",
                        "allowed_roles": ["employee", "hr"],
                    }
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    index_manifest = tmp_path / "index_manifest.json"

    monkeypatch.setattr(kb, "_require_usable_tempdir", lambda: None)
    monkeypatch.setattr(kb, "MANIFEST_PATH", manifest)
    monkeypatch.setattr(kb, "KNOWLEDGE_DIR", knowledge)
    monkeypatch.setattr(kb, "INDEX_MANIFEST_PATH", index_manifest)
    monkeypatch.setattr(kb, "VECTOR_STORE_DIR", tmp_path / "store")
    monkeypatch.setattr(kb, "_build_embeddings", lambda: (object(), "fake-model"))
    monkeypatch.setattr(kb, "_warn_embedding_model_changed", lambda _m: None)
    monkeypatch.setattr(kb, "_reset_collection", lambda: None)
    monkeypatch.setattr(
        kb.Chroma, "from_documents", staticmethod(lambda **kwargs: None)
    )

    summary = kb.build_vector_store()

    assert summary["document_count"] == 1
    assert summary["chunk_count"] >= 1
    assert index_manifest.exists()
    written = json.loads(index_manifest.read_text(encoding="utf-8"))
    assert written["embedding_model"] == "fake-model"
